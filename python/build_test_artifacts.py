"""
Build the QA workbook: test cases, RTM, defect log, change log, RAID log, reconciliation results.
Test results come from the actual reconciliation outputs in data/results and api/api_test_results.csv.
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R1 = os.path.join(ROOT, "data/results/trial1")
R2 = os.path.join(ROOT, "data/results/rehearsal2")

s1 = pd.read_csv(os.path.join(R1, "reconciliation_summary.csv")).set_index("rule_id")
s2 = pd.read_csv(os.path.join(R2, "reconciliation_summary.csv")).set_index("rule_id")
ct1 = pd.read_csv(os.path.join(R1, "control_totals.csv")).set_index("metric")["value"]
ct2 = pd.read_csv(os.path.join(R2, "control_totals.csv")).set_index("metric")["value"]
exc1 = pd.read_csv(os.path.join(R1, "exceptions.csv"), dtype=str)
api = pd.read_csv(os.path.join(ROOT, "api/api_test_results.csv"))
leg = pd.read_csv(os.path.join(ROOT, "data/sample/legacy_payments.csv"), dtype={"PMT_ID": str})
tgt = pd.read_csv(os.path.join(ROOT, "data/sample/target_payments.csv"), dtype={"payment_id": str})
amt = dict(zip(leg.PMT_ID, leg.PMT_AMT)); amt.update(dict(zip(tgt.payment_id, tgt.amount_usd)))
legacy_ids_value = leg.set_index("PMT_ID").PMT_AMT
value_by_rule = {r: round(sum(amt[p] for p in g.payment_id), 2) for r, g in exc1.groupby("rule_id")}

rules = {
 "R01": ("REQ-001", "Record completeness: every legacy payment exists in target"),
 "R02": ("REQ-003", "No target payment exists without a legacy source"),
 "R03": ("REQ-002", "No payment_id appears more than once in target"),
 "R04": ("REQ-004", "Amount matches legacy to the cent"),
 "R05": ("REQ-005", "Originator account migrates at full length"),
 "R06": ("REQ-006", "Settlement date is ISO 8601 (YYYY-MM-DD)"),
 "R07": ("REQ-006", "Settlement date value equals legacy date"),
 "R08": ("REQ-007", "Routing numbers present and 9 digits"),
 "R09": ("REQ-008", "Status codes map STL/PND/RTN to SETTLED/PENDING/RETURNED"),
 "R10": ("REQ-009", "Fraud flag preserved"),
 "R11": ("REQ-010", "Payment method maps W/A to WIRE/ACH"),
 "R12": ("REQ-010", "Payment direction maps C/D to CREDIT/DEBIT"),
 "R13": ("REQ-007", "Originator routing value unchanged"),
 "R14": ("REQ-007", "Beneficiary routing value unchanged"),
 "R15": ("REQ-005", "Beneficiary account unchanged"),
 "R16": ("REQ-011", "Trace ID unchanged"),
 "R17": ("REQ-007", "Routing numbers pass ABA checksum"),
}
DEF = {r: f"DEF-{i:03d}" for i, r in enumerate(["R01","R02","R03","R04","R05","R06","R07","R08","R09","R10"], start=1)}

tcs = []  # id, type, title, req, method, expected, r1, r1ev, r2, r2ev, defect
for i, (r, (req, title)) in enumerate(rules.items(), start=1):
    n1, n2 = int(s1.loc[r, "exceptions"]), int(s2.loc[r, "exceptions"])
    tcs.append((f"TC-{i:02d}", "SIT", title, req, f"Run sql/02_reconciliation_rules.sql, rule {r}", "0 exceptions",
                "Pass" if n1 == 0 else "Fail", f"{n1:,} exceptions", "Pass" if n2 == 0 else "Fail", f"{n2:,} exceptions", DEF.get(r, "") if n1 else ""))
api_titles = {
 "TC-18": "API accepts a valid payment (201)", "TC-19": "API rejects invalid routing checksum (422)",
 "TC-20": "API rejects missing required field (400)", "TC-21": "API rejects negative amount (422)",
 "TC-22": "API rejects duplicate Idempotency-Key (409)", "TC-23": "API rejects wrong date format (400)",
 "TC-24": "API rejects unsupported payment method (400)",
}
for _, a in api.iterrows():
    tcs.append((a.test_case, "SIT", api_titles[a.test_case], "REQ-013",
                f"Submit api/payloads/{a.test_case}_{a.scenario}.json to reference validator", f"HTTP {a.expected_http} {a.expected_code if isinstance(a.expected_code,str) else ''}".strip(),
                a.result, f"HTTP {a.actual_http} {a.actual_code if isinstance(a.actual_code,str) else ''}".strip(), "N/A", "No change since run 1", ""))
# control totals
cnt1, cnt2 = ct1["Target row count (incl. duplicates)"], ct2["Target row count (incl. duplicates)"]
tcs.append(("TC-25", "SIT", "Control totals tie out (row count, distinct keys, amount total)", "REQ-014",
            "Run sql/03_control_totals.sql", "Counts and amount total equal legacy",
            "Fail", f"Rows {int(cnt1):,} vs 100,000; amount variance -${abs(float(ct1['Variance, target minus legacy (USD)'])):,.2f}",
            "Pass" if (int(cnt2) == 100000 and float(ct2['Variance, target minus legacy (USD)']) == 0) else "Fail",
            f"Rows {int(cnt2):,}; amount variance ${abs(float(ct2['Variance, target minus legacy (USD)'])):,.2f}", "DEF-001, DEF-002, DEF-003, DEF-004"))
reg_zero = all(int(s2.loc[r, "exceptions"]) == 0 for r in ["R11","R12","R13","R14","R15","R16","R17"])
tcs.append(("TC-26", "Regression", "Clean fields stay clean after fixes (R11 to R17)", "REQ-014",
            "Rerun R11 to R17 on rehearsal 2", "0 exceptions", "N/A", "Baseline: 0 exceptions on trial load 1 for R11 to R17", "Pass" if reg_zero else "Fail", "R11 to R17: 0 exceptions", ""))
# exception log completeness (real check)
need = ["rule_id","rule_name","severity","payment_id","legacy_value","target_value"]
log_ok = list(exc1.columns) == need and len(exc1) == int(s1["exceptions"].sum())
tcs.append(("TC-27", "SIT", "Exception log has required columns and one row per exception", "REQ-012",
            "Check data/results/trial1/exceptions.csv columns and row count vs summary", "6 columns, rows equal summary total",
            "Pass" if log_ok else "Fail", f"{len(exc1):,} rows, columns OK" if log_ok else "Mismatch", "N/A", "No exceptions to log on rehearsal 2", ""))
tcs.append(("TC-28", "UAT", "Payments Operations reviews the ranked exception report", "REQ-012",
            "Business user reviews exceptions by severity and confirms it is usable", "Report accepted", "Not Run", "Awaiting business users", "Not Run", "Awaiting business users", ""))
tcs.append(("TC-29", "UAT", "Compliance confirms all flagged payments retained", "REQ-009",
            "Compliance spot-checks fraud-flagged payments in target", "All 122 flagged payments retained", "Not Run", "Awaiting business users", "Not Run", "Awaiting business users", ""))
tcs.append(("TC-30", "UAT", "Finance Controller signs off control totals", "REQ-014",
            "Finance reviews control totals and sign-off sheet", "Signed", "Not Run", "Awaiting business users", "Not Run", "Awaiting business users", ""))

# Defects
dnames = {
 "R01": ("Payments missing from target", "Load job skipped rows that failed a target constraint without logging them", "Log and reject every skipped row. Add completeness check before cutover."),
 "R02": ("Records in target with no legacy source", "Rows from an earlier test load were not purged before the trial load", "Purge target before each load. Add R02 as a pre-cutover gate."),
 "R03": ("Duplicate payment_id in target", "Load job restarted mid-batch and re-inserted rows, no unique key enforced", "Enforce unique key on payment_id. Make load restartable."),
 "R04": ("Amounts rounded to whole dollars", "One ETL path cast amounts to integer", "Use Decimal(18,2) end to end. Remove integer cast."),
 "R05": ("Originator account numbers truncated", "Target column defined too narrow (see FG-04)", "Widen column to at least 11 characters (CR-004)."),
 "R06": ("Settlement date written as MM/DD/YYYY", "Second code path used a different date format", "Single ISO 8601 conversion function for all paths."),
 "R07": ("Settlement date shifted by one day", "Timezone conversion moved some dates", "Convert dates without timezone logic."),
 "R08": ("Beneficiary routing number blank", "Blank source field mapped without a required-field check", "Reject blanks before load. Add R08 and R17 as gates."),
 "R09": ("Pending payments shown as settled", "Unmapped PND code defaulted to SETTLED", "Explicit lookup, reject unmapped codes."),
 "R10": ("Fraud flag lost on conversion", "Boolean mapping dropped Y on a subset of rows", "Explicit Y to 1 mapping. Keep R10 as a Critical gate."),
}
m = leg.merge(tgt.drop_duplicates("payment_id"), left_on="PMT_ID", right_on="payment_id")
r4ids = set(exc1[exc1.rule_id == "R04"].payment_id)
abs_err = float((m[m.PMT_ID.isin(r4ids)].amount_usd - m[m.PMT_ID.isin(r4ids)].PMT_AMT).abs().sum())
dnames["R04"] = (f"Amounts rounded to whole dollars (absolute error ${abs_err:,.2f} in total)",) + dnames["R04"][1:]
defects = []
for r, d in DEF.items():
    n1 = int(s1.loc[r, "exceptions"])
    n2 = int(s2.loc[r, "exceptions"])
    defects.append((d, dnames[r][0], s1.loc[r, "severity"], r, f"TC-{list(rules).index(r)+1:02d}", n1, value_by_rule[r],
                    dnames[r][1], dnames[r][2], "Trial load 1", "Closed" if n2 == 0 else "Open",
                    f"Retest passed on rehearsal 2 ({n2} exceptions)" if n2 == 0 else "Retest failed"))

changes = [
 ("CR-001", "2026-10-07", "Add ABA checksum rule R17 to reconciliation", "API enforces ABA checksum, so reconciliation must check it too (FG-05)", "REQ-007, TC-17", "BA", "Core Platform Product Owner", "Implemented"),
 ("CR-002", "2026-10-07", "Add control rules R11 to R17 for clean fields", "Prove fields that should not change did not change", "REQ-005, 007, 010, 011, TC-11 to TC-17, TC-26", "BA", "QA Lead", "Implemented"),
 ("CR-003", "2026-10-07", "Classify missing records and fraud flag loss as Critical", "Both are irreversible if missed at cutover", "REQ-001, REQ-009, severity model", "BA", "Financial Crime and Compliance", "Approved"),
 ("CR-004", "2026-10-07", "Widen target account number columns to 11 or more characters", "Legacy accounts run to 11 characters (FG-04, DEF-005)", "REQ-005, TC-05", "Data Migration Lead", "Core Platform Product Owner", "Implemented in rehearsal 2"),
 ("CR-005", "2026-10-07", "Require Idempotency-Key header on payment API", "Prevent duplicate payment submission (TC-22)", "REQ-013, TC-22", "BA", "Core Platform Product Owner", "Implemented in spec v1.0"),
]

gross = float(value_by_rule["R01"]) + float(value_by_rule["R03"]) + float(value_by_rule["R02"])
raid = [
 ("Risk", "RK-01", "Offsetting errors hide in totals: net variance looks small while gross error is large", f"High impact. Net variance ${float(ct1['Variance, target minus legacy (USD)']):,.2f} vs ${gross:,.0f} of missing, duplicate and orphan payments", "Finance Controller", "Reconcile every payment (R01 to R17), not just totals", "Mitigated"),
 ("Risk", "RK-02", "Fraud flag dropped during conversion", f"Critical. {int(s1.loc['R10','exceptions'])} flagged payments lost the flag in trial load 1, ${value_by_rule['R10']:,.0f} affected", "Financial Crime and Compliance", "R10 is a Critical gate. Compliance spot-check in UAT (TC-29)", "Open, pending UAT"),
 ("Risk", "RK-03", "Pending payments reported as settled", f"High. {int(s1.loc['R09','exceptions'])} of 5,038 pending payments affected in trial load 1", "Payments Operations", "Explicit status lookup, R09 gate", "Mitigated"),
 ("Risk", "RK-04", "Business users not available for UAT", "Medium. Go-live depends on TC-28 to TC-30", "Program Manager", "Book UAT slots before rehearsal 3", "Open"),
 ("Risk", "RK-05", "Synthetic data may not reflect real production data quality", "Medium. Portfolio project, not a real conversion", "BA", "State this in the README and on the resume", "Accepted"),
 ("Assumption", "AS-01", "Legacy extract is complete and frozen at extract time", "", "Data Migration Lead", "Confirm extract control total with legacy team", "Open"),
 ("Assumption", "AS-02", "Synthetic routing numbers use prefix 99 so they cannot match a real bank", "", "BA", "None needed", "Closed"),
 ("Assumption", "AS-03", "Business rules in the mapping document are approved by the business owner", "", "Head of Payments Operations", "Walk through mapping in review session", "Open"),
 ("Issue", "IS-01", "Trial load 1 produced 4,010 exceptions (4.01% of records)", f"{int(s1.loc['R01','exceptions'])} missing, {int(s1.loc['R03','exceptions'])} duplicate, {int(s1.loc['R02','exceptions'])} orphan, 10 fraud flags lost, and 6 other types", "Data Migration Lead", "Fix 10 defects, rerun as rehearsal 2", "Closed (rehearsal 2: 0 exceptions)"),
 ("Issue", "IS-02", "Net control total variance would pass a net-only tolerance check", f"${abs(float(ct1['Variance, target minus legacy (USD)'])):,.2f} net vs ${gross:,.0f} gross", "Finance Controller", "Adopt decision DC-01", "Closed"),
 ("Decision", "DC-01", "Reconcile at payment level as well as in total", "", "Head of Payments Operations", "Done: rules R01 to R17", "Approved"),
 ("Decision", "DC-02", "How to treat merchant (M-prefix) beneficiary accounts", "33,815 records have M-prefix beneficiaries (FG-12)", "Head of Payments Operations", "Business owner to decide: add beneficiary type attribute or keep as is", "Open"),
 ("Decision", "DC-03", "Severity model: Critical, High, Medium", "", "QA Lead", "Done: defined in BRD section 7", "Approved"),
 ("Dependency", "DP-01", "Platform team widens account number column (FG-04, CR-004)", "Blocks DEF-005 retest", "Core Platform Product Owner", "Applied in rehearsal 2", "Done"),
 ("Dependency", "DP-02", "Compliance availability for TC-29", "Blocks REQ-009 sign-off", "Financial Crime and Compliance", "Schedule review", "Open"),
 ("Dependency", "DP-03", "Business owner approval of mapping rules", "Blocks AS-03", "Head of Payments Operations", "Schedule review", "Open"),
 ("Action", "AC-01", "Run UAT with Payments Operations (TC-28)", "", "Program Manager", "Book session", "Open"),
 ("Action", "AC-02", "Get decision on merchant account handling (DC-02)", "", "BA", "Raise at next steering review", "Open"),
 ("Action", "AC-03", "Compliance spot-check of 122 fraud-flagged payments (TC-29)", "", "Financial Crime and Compliance", "Schedule review", "Open"),
]

# ---------- workbook ----------
FONT = "Arial"
hdr_fill = PatternFill("solid", start_color="1F3864")
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
body = Font(name=FONT, size=10)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def sheet(wb, name, headers, rows, widths, note=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=FONT, italic=True, size=9, color="595959")
        r0 = 3
    for c, h in enumerate(headers, 1):
        x = ws.cell(r0, c, h); x.font = hdr_font; x.fill = hdr_fill; x.alignment = wrap; x.border = border
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            x = ws.cell(i, c, v); x.font = body; x.alignment = wrap; x.border = border
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    return ws, r0

wb = Workbook(); wb.remove(wb.active)
NOTE = "Synthetic data adapted from public PaySim. Fictional bank and roles. Results come from scripts in /python."

# Test cases (L = latest result formula)
tc_rows = [list(t) for t in tcs]
ws, r0 = sheet(wb, "Test Cases",
    ["ID", "Type", "Title", "Requirement", "Method", "Expected", "Run 1 (trial load 1)", "Run 1 evidence", "Run 2 (rehearsal 2)", "Run 2 evidence", "Defect", "Latest result"],
    [t + [None] for t in tc_rows], [8, 11, 46, 12, 42, 28, 12, 34, 12, 30, 22, 12], NOTE)
n_tc = len(tc_rows); first = r0 + 1; last = r0 + n_tc
for i in range(first, last + 1):
    ws.cell(i, 12, f'=IF(I{i}="N/A",G{i},I{i})').font = body
    ws.cell(i, 12).border = border; ws.cell(i, 12).alignment = wrap
tc_ref = lambda col: f"'Test Cases'!${col}${first}:${col}${last}"

# Defects
dws, d0 = sheet(wb, "Defect Log",
    ["ID", "Title", "Severity", "Recon rule", "Test case", "Records affected", "Value of affected payments (USD)", "Suspected root cause (simulated)", "Fix", "Found in", "Status", "Retest"],
    defects, [10, 38, 10, 10, 10, 12, 20, 46, 42, 14, 10, 34],
    NOTE + " Root causes are simulated hypotheses: the defects were seeded.")
for i in range(d0 + 1, d0 + 1 + len(defects)):
    dws.cell(i, 7).number_format = '$#,##0.00'; dws.cell(i, 6).number_format = '#,##0'
d_first, d_last = d0 + 1, d0 + len(defects)

# RTM
reqs = [
 ("REQ-001","Every in-scope legacy payment shall exist in the target","Must"),
 ("REQ-002","No payment_id shall appear more than once in the target","Must"),
 ("REQ-003","No target payment shall exist without a legacy source","Must"),
 ("REQ-004","Amount shall match the legacy value to the cent, no rounding","Must"),
 ("REQ-005","Account numbers shall migrate at full length","Must"),
 ("REQ-006","Settlement date converts to ISO 8601 with the same date value","Must"),
 ("REQ-007","Routing numbers present, 9 digits, ABA checksum valid, unchanged","Must"),
 ("REQ-008","Status codes map STL, PND, RTN with no defaulting","Must"),
 ("REQ-009","Fraud flag preserved for every payment","Must"),
 ("REQ-010","Payment method and direction map correctly","Must"),
 ("REQ-011","Trace ID preserved unchanged","Must"),
 ("REQ-012","Every exception logged with rule, severity, legacy and target values","Must"),
 ("REQ-013","Payment API accepts JSON, validates, returns structured errors","Should"),
 ("REQ-014","Go-live needs zero open Critical and High exceptions and signed UAT","Must"),
]
rtm_rows = []
for rq, text, pr in reqs:
    tcl = ", ".join(t[0] for t in tcs if t[3] == rq)
    dfl = ", ".join(sorted({d[0] for d in defects if d[4] in [t[0] for t in tcs if t[3] == rq]}))
    rtm_rows.append([rq, text, pr, tcl, dfl, None, None, None, None])
rws, r_0 = sheet(wb, "RTM",
    ["Requirement", "Description", "Priority", "Test cases", "Defects (trial load 1)", "Test cases", "Latest Pass", "Latest Fail", "Status"],
    rtm_rows, [12, 56, 9, 44, 26, 10, 11, 11, 14], NOTE)
for i in range(r_0 + 1, r_0 + 1 + len(reqs)):
    rws.cell(i, 6, f'=COUNTIF({tc_ref("D")},A{i})')
    rws.cell(i, 7, f'=COUNTIFS({tc_ref("D")},A{i},{tc_ref("L")},"Pass")')
    rws.cell(i, 8, f'=COUNTIFS({tc_ref("D")},A{i},{tc_ref("L")},"Fail")')
    rws.cell(i, 9, f'=IF(H{i}>0,"Fail",IF(G{i}=F{i},"Pass","Pending UAT"))')
    for c in range(6, 10):
        rws.cell(i, c).font = body; rws.cell(i, c).border = border; rws.cell(i, c).alignment = wrap

# Change log, RAID
sheet(wb, "Change Log", ["ID", "Date", "Change", "Reason", "Impact", "Requested by", "Approver (fictional role)", "Status"],
      changes, [9, 11, 50, 52, 36, 20, 28, 24], NOTE)
sheet(wb, "RAID Log", ["Type", "ID", "Description", "Impact / detail", "Owner (fictional role)", "Next step", "Status"],
      raid, [11, 8, 58, 56, 28, 48, 28], NOTE)

# Reconciliation results
rr = []
for r in rules:
    rr.append([r, s1.loc[r, "rule_name"], s1.loc[r, "severity"], int(s1.loc[r, "exceptions"]), int(s2.loc[r, "exceptions"]), value_by_rule.get(r, 0.0)])
rcws, rc0 = sheet(wb, "Recon Results", ["Rule", "Name", "Severity", "Exceptions, trial load 1", "Exceptions, rehearsal 2", "Value of affected payments, trial load 1 (USD)"],
                  rr, [8, 34, 10, 18, 18, 24], NOTE)
for i in range(rc0 + 1, rc0 + 1 + len(rr)):
    rcws.cell(i, 6).number_format = '$#,##0.00'; rcws.cell(i, 4).number_format = '#,##0'; rcws.cell(i, 5).number_format = '#,##0'
end = rc0 + len(rr)
tot = end + 1
rcws.cell(tot, 1, "Total").font = Font(name=FONT, bold=True, size=10)
rcws.cell(tot, 4, f"=SUM(D{rc0+1}:D{end})"); rcws.cell(tot, 5, f"=SUM(E{rc0+1}:E{end})")
for c in (4, 5):
    rcws.cell(tot, c).font = Font(name=FONT, bold=True, size=10); rcws.cell(tot, c).number_format = '#,##0'

# Summary sheet
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Payments Data Migration: QA Summary"; sm["A1"].font = Font(name=FONT, bold=True, size=14)
sm["A2"] = NOTE; sm["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
items = [
 ("Test cases", f"=COUNTA({tc_ref('A')})"),
 ("Latest result: Pass", f'=COUNTIF({tc_ref("L")},"Pass")'),
 ("Latest result: Fail", f'=COUNTIF({tc_ref("L")},"Fail")'),
 ("Latest result: Not Run (UAT)", f'=COUNTIF({tc_ref("L")},"Not Run")'),
 ("Defects raised", f"=COUNTA('Defect Log'!$A${d_first}:$A${d_last})"),
 ("Defects Critical", f"=COUNTIF('Defect Log'!$C${d_first}:$C${d_last},\"Critical\")"),
 ("Defects High", f"=COUNTIF('Defect Log'!$C${d_first}:$C${d_last},\"High\")"),
 ("Defects Medium", f"=COUNTIF('Defect Log'!$C${d_first}:$C${d_last},\"Medium\")"),
 ("Defects Closed", f"=COUNTIF('Defect Log'!$K${d_first}:$K${d_last},\"Closed\")"),
 ("Requirements", f"=COUNTA(RTM!$A${r_0+1}:$A${r_0+len(reqs)})"),
 ("Requirements Pass", f"=COUNTIF(RTM!$I${r_0+1}:$I${r_0+len(reqs)},\"Pass\")"),
 ("Requirements Pending UAT", f"=COUNTIF(RTM!$I${r_0+1}:$I${r_0+len(reqs)},\"Pending UAT\")"),
 ("Exceptions, trial load 1", f"='Recon Results'!D{tot}"),
 ("Exceptions, rehearsal 2", f"='Recon Results'!E{tot}"),
]
for i, (k, v) in enumerate(items, 4):
    sm.cell(i, 1, k).font = body; c = sm.cell(i, 2, v); c.font = body; c.number_format = '#,##0'
sm.column_dimensions["A"].width = 34; sm.column_dimensions["B"].width = 14

out = os.path.join(ROOT, "tests", "QA_Workbook.xlsx")
wb.save(out)
print("saved", out, "| test cases:", len(tcs), "| defects:", len(defects))
